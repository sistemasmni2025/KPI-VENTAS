import polars as pl
import os
import PyPDF2
import docx
import pandas as pd

def analizar_archivo_datos(ruta_archivo: str, hoja: str = None) -> str:
    """
    Lee un archivo local (Excel, CSV, Word, PDF, ODS, TXT) y extrae su contenido.
    Soporta selección de hojas para Excel y es robusto ante errores de tipos de datos.
    """
    if not os.path.exists(ruta_archivo):
        return f"Error: No se encontró el archivo en la ruta {ruta_archivo}"
        
    ext = ruta_archivo.split('.')[-1].lower()
    
    try:
        if ext in ['csv', 'xlsx', 'xls', 'ods']:
            lista_hojas = []
            if ext in ['xlsx', 'xls', 'ods']:
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(ruta_archivo, read_only=True)
                    lista_hojas = wb.sheetnames
                except:
                    pass

            # LEER DATOS CON PANDAS PARA MÁXIMA ROBUSTEZ EN TIPOS DE DATOS
            # Usamos dtype=str para evitar que falle por inferencia de esquema (infer_schema_length)
            if ext == 'csv':
                df_pd = pd.read_csv(ruta_archivo, dtype=str, on_bad_lines='skip')
            else:
                try:
                    if hoja:
                        df_pd = pd.read_excel(ruta_archivo, sheet_name=hoja, dtype=str)
                    else:
                        df_pd = pd.read_excel(ruta_archivo, dtype=str) # Primera hoja
                except Exception as e_hoja:
                    return f"Error al leer la hoja '{hoja}': {str(e_hoja)}. Hojas disponibles: {lista_hojas}"
            
            # Limpiar nombres de columnas originales para el mapeo
            columnas_originales = list(df_pd.columns)
            nuevas_columnas = [f"columna_{i}" for i in range(len(df_pd.columns))]
            df_pd.columns = nuevas_columnas
            
            # Crear mapa de columnas
            mapeo_columnas = "\n".join([f"- Original '{orig}' -> Usar en SQL: `columna_{i}`" for i, orig in enumerate(columnas_originales)])
            
            filas = len(df_pd)
            
            # Subir a la BD de PostgreSQL temporalmente
            db_uri = f"postgresql://{os.getenv('DB_USER')}:{os.getenv('DB_PASSWORD')}@{os.getenv('DB_HOST')}:{os.getenv('DB_PORT')}/{os.getenv('DB_NAME')}"
            
            try:
                import sqlalchemy
                engine = sqlalchemy.create_engine(db_uri)
                # Cargamos todo como TEXT en Postgres para que orbis_safe_numeric haga el trabajo sucio
                df_pd.to_sql("temp_archivo_usuario", engine, if_exists='replace', index=False)
                
                # Muestra de 5 filas reales
                muestra = df_pd.head(5).to_markdown(index=False)
                
                info_hojas = f"Hojas en Excel: {lista_hojas}\n" if lista_hojas else ""
                hoja_actual = f"Hoja procesada: '{hoja if hoja else lista_hojas[0] if lista_hojas else 'N/A'}'\n"
                
                resumen = (
                    f"=== DATOS DEL ARCHIVO {ext.upper()} ===\n"
                    f"{info_hojas}{hoja_actual}"
                    f"Total de filas: {filas}\n\n"
                    f"MAPEO DE COLUMNAS (ASIGNACIÓN EXACTA):\n{mapeo_columnas}\n\n"
                    f"Muestra de 5 filas reales:\n{muestra}\n\n"
                    f"[SISTEMA A IA: Archivo cargado en la tabla `temp_archivo_usuario`. "
                    f"REGLA: Usa exclusivamente `columna_0`, `columna_1`, etc. "
                    f"Usa siempre `orbis_safe_numeric(columna_X)` para SUM y filtros numéricos. "
                    f"Esto es CRÍTICO porque los datos se cargaron como texto para evitar errores de parseo.]\n"
                )
                return resumen
            except Exception as e_bd:
                muestra = df_pd.head(15).to_markdown(index=False)
                return f"Error en Sandbox BD ({str(e_bd)}). Muestra:\n{muestra}"
            
        elif ext == 'pdf':
            reader = PyPDF2.PdfReader(ruta_archivo)
            texto = "".join([p.extract_text() + "\n" for p in reader.pages[:5]])
            return f"=== CONTENIDO PDF ===\n\n{texto[:5000]}"
            
        elif ext in ['doc', 'docx']:
            doc = docx.Document(ruta_archivo)
            texto = "\n".join([para.text for para in doc.paragraphs])
            return f"=== CONTENIDO WORD ===\n\n{texto[:5000]}"
            
        elif ext in ['txt', 'md', 'json', 'py', 'js', 'jsx', 'html', 'css']:
            with open(ruta_archivo, 'r', encoding='utf-8', errors='ignore') as f:
                return f"=== CONTENIDO TXT ===\n\n{f.read()[:5000]}"
            
        else:
            return f"Error: Formato .{ext} no soportado."
            
    except Exception as e:
        return f"Error general: {str(e)}"
